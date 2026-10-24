# -*- coding: utf-8 -*-
"""
MOTOR DE VISION ARTIFICIAL - Taller 5 (inspeccion de contenedores tapados)
==========================================================================
Vision artificial clasica (OpenCV), sin deep learning:
  1. Segmentacion por color HSV (tapa = azul, contenedor = blanco).
  2. Medicion de: inclinacion de la tapa, altura visible de la tapa (si la tapa
     esta sellada solo se ve su borde; si falta presion se ve la pestana
     interior y la altura aumenta) y desalineacion lateral tapa/contenedor.
  3. Reglas de decision explicables (justificacion tecnica por imagen).
  4. Calibracion asistida: aprende los umbrales a partir de las respuestas del
     operador (APTA / RECHAZADA) con regularizacion para evitar sobreajuste.
  5. Registro historico en Excel (acumula sin sobrescribir).
Este modulo no tiene interfaz; lo usa app.py (Streamlit).
"""
import itertools
import json
import time
import unicodedata
from datetime import datetime
from pathlib import Path

import cv2
import numpy as np
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# ---------------------------------------------------------------------------
# PARAMETROS (las imagenes se normalizan a 752 x 480 px antes de medir)
# ---------------------------------------------------------------------------
ANCHO, ALTO = 752, 480
ROI_X_MIN, ROI_Y_MAX = 45, 446       # excluye la interfaz de la camara
MIN_AREA_CONTENEDOR = 5000           # px blancos minimos  -> hay contenedor
MIN_AREA_TAPA = 5000                 # px azules minimos   -> hay tapa

UMBRALES_DEFECTO = {
    "inclinacion": 4.0,      # grados maximos de inclinacion de la tapa
    "altura": 62.0,          # px maximos de tapa visible (sellada ~50 px)
    "desalineacion": 45.0,   # px maximos de desplazamiento lateral
    # segmentacion HSV de OpenCV (H 0-179, S y V 0-255)
    "azul_s_min": 120, "azul_v_min": 60,
    "blanco_s_max": 70, "blanco_v_min": 110,
}
CLAVES_DECISION = ("inclinacion", "altura", "desalineacion")
LIMITES = {"inclinacion": (1.0, 10.0), "altura": (40.0, 120.0),
           "desalineacion": (15.0, 80.0)}      # rangos fisicamente plausibles
RUTA_CALIBRACION = Path(__file__).resolve().parent / "calibracion.json"

# Causas de rechazo y accion de disposicion sugerida (vista MES)
C_SIN_CONT = "Sin contenedor"
C_SIN_TAPA = "Sin tapa"
C_INCL = "Tapa inclinada"
C_SELLADO = "Tapa sin sellar (falta presión)"
C_DESAL = "Tapa desalineada"
ACCIONES = {
    C_SIN_CONT: "Verificar alimentación y sensor de presencia",
    C_SIN_TAPA: "Reprocesar: colocar tapa",
    C_INCL: "Reprocesar: reasentar la tapa",
    C_SELLADO: "Reprocesar: aplicar presión de cierre",
    C_DESAL: "Reprocesar: centrar y presionar la tapa",
}

COLUMNAS = ["Lote", "Imagen", "Fecha", "Hora", "Contenedor", "Tapa",
            "Inclinación (°)", "Altura (px)", "Desalineación (px)",
            "Resultado", "Justificación Técnica", "Causa de rechazo"]


def cargar_calibracion(ruta=RUTA_CALIBRACION):
    ruta = Path(ruta)
    if not ruta.exists():
        return None
    try:
        return json.loads(ruta.read_text(encoding="utf-8"))
    except (OSError, ValueError):
        return None


UMBRALES = dict(UMBRALES_DEFECTO)
_cal = cargar_calibracion()
if _cal:
    UMBRALES.update(_cal.get("umbrales", {}))


# ---------------------------------------------------------------------------
# Lectura y segmentacion
# ---------------------------------------------------------------------------
def decodificar(datos):
    """bytes de un archivo de imagen -> imagen BGR normalizada."""
    img = cv2.imdecode(np.frombuffer(datos, dtype=np.uint8), cv2.IMREAD_COLOR)
    if img is None:
        raise ValueError("no se pudo leer la imagen")
    return cv2.resize(img, (ANCHO, ALTO))


def leer_imagen(ruta):
    return decodificar(Path(ruta).read_bytes())


def _mayor_componente(mascara):
    n, etiquetas, stats, _ = cv2.connectedComponentsWithStats(mascara, 8)
    if n < 2:
        return np.zeros_like(mascara)
    k = 1 + int(np.argmax(stats[1:, cv2.CC_STAT_AREA]))
    return (etiquetas == k).astype(np.uint8)


def _mascara_azul(hsv, u):
    m = ((hsv[..., 0] > 100) & (hsv[..., 0] < 135) &
         (hsv[..., 1] > u["azul_s_min"]) & (hsv[..., 2] > u["azul_v_min"])).astype(np.uint8)
    m[:, :ROI_X_MIN] = 0
    m[ROI_Y_MAX:] = 0
    return _mayor_componente(cv2.morphologyEx(m, cv2.MORPH_OPEN, np.ones((5, 5), np.uint8)))


def _mascara_blanco(hsv, u):
    m = ((hsv[..., 1] < u["blanco_s_max"]) & (hsv[..., 2] > u["blanco_v_min"])).astype(np.uint8)
    m[:, :ROI_X_MIN] = 0
    m[ROI_Y_MAX:] = 0
    m[:60] = 0            # textos de la interfaz de la camara
    m[:, 735:] = 0        # boton lateral de la interfaz
    return _mayor_componente(cv2.morphologyEx(m, cv2.MORPH_OPEN, np.ones((5, 5), np.uint8)))


def _hsv(img):
    """Prefiltro gaussiano 5x5 (reduce el ruido del sensor) y conversion a HSV."""
    return cv2.cvtColor(cv2.GaussianBlur(img, (5, 5), 0), cv2.COLOR_BGR2HSV)


def segmentar(img, u=None):
    u = {**UMBRALES, **(u or {})}
    hsv = _hsv(img)
    return _mascara_azul(hsv, u), _mascara_blanco(hsv, u)


# ---------------------------------------------------------------------------
# Medicion y decision
# ---------------------------------------------------------------------------
def _medir_mascaras(azul, blanco):
    f = {"contenedor": int(blanco.sum()) > MIN_AREA_CONTENEDOR,
         "tapa": int(azul.sum()) > MIN_AREA_TAPA,
         "inclinacion": None, "altura": None, "desalineacion": None,
         "caja_contenedor": None, "caja_tapa": None, "linea_tapa": None}
    if f["contenedor"]:
        f["caja_contenedor"] = cv2.boundingRect(blanco)
    if f["tapa"]:
        f["caja_tapa"] = cv2.boundingRect(azul)
    if f["contenedor"] and f["tapa"]:
        ys, xs = np.nonzero(azul)
        pts = np.column_stack([xs, ys]).astype(np.float32)
        vx, vy, cx, cy = cv2.fitLine(pts[::7], cv2.DIST_L2, 0, 0.01, 0.01).ravel()
        inc = abs(float(np.degrees(np.arctan2(vy, vx))))
        inc = min(inc, 180 - inc)
        cols = range(int(xs.min()) + 60, int(xs.max()) - 60, 10)
        altura = float(np.median([np.ptp(np.nonzero(azul[:, c])[0])
                                  for c in cols if azul[:, c].any()]))
        xb = np.nonzero(blanco)[1]
        f.update(inclinacion=round(inc, 1), altura=round(altura, 1),
                 desalineacion=round(float(abs(xs.mean() - xb.mean())), 1),
                 linea_tapa=(float(cx), float(cy), float(vx), float(vy)))
    return f


def medir(img, u=None):
    """Caracteristicas geometricas de la imagen (sin decidir)."""
    return _medir_mascaras(*segmentar(img, u))


def decidir(m, u=None):
    """Reglas de decision. Devuelve (resultado, causas, justificacion)."""
    u = {**UMBRALES, **(u or {})}
    if not m["contenedor"]:
        return ("RECHAZADO", [C_SIN_CONT],
                "No se detecta contenedor en la imagen (campo oscuro).")
    if not m["tapa"]:
        return ("RECHAZADO", [C_SIN_TAPA],
                "Contenedor detectado pero sin tapa azul: contenedor destapado.")
    inc, alt, des = m["inclinacion"], m["altura"], m["desalineacion"]
    causas, motivos = [], []
    if inc > u["inclinacion"]:
        causas.append(C_INCL)
        motivos.append(f"tapa inclinada {inc:.1f}° (máx. {u['inclinacion']:.1f}°): "
                       f"no apoya de forma uniforme sobre el borde")
    if des > u["desalineacion"]:
        causas.append(C_DESAL)
        motivos.append(f"tapa descentrada {des:.0f} px respecto al contenedor "
                       f"(máx. {u['desalineacion']:.0f} px): no podrá cerrar")
    # La altura de sellado solo es significativa si la tapa esta horizontal y
    # centrada; por eso se evalua en ultimo lugar (decision jerarquica).
    if not causas and alt > u["altura"]:
        causas.append(C_SELLADO)
        motivos.append(f"tapa sin sellar por completo: se ve la pestaña interior "
                       f"(altura visible {alt:.0f} px > máx. {u['altura']:.0f} px), "
                       f"falta presión de cierre")
    if causas:
        texto = "; ".join(motivos)
        return "RECHAZADO", causas, texto[0].upper() + texto[1:] + "."
    return ("APROBADO", [],
            f"Tapa presente, horizontal ({inc:.1f}° ≤ {u['inclinacion']:.1f}°), sellada "
            f"(altura {alt:.0f} px ≤ {u['altura']:.0f} px) y alineada "
            f"({des:.0f} px ≤ {u['desalineacion']:.0f} px).")


def _ascii(texto):
    texto = texto.replace("°", " deg")
    return unicodedata.normalize("NFKD", texto).encode("ascii", "ignore").decode()


def dibujar(img, m, resultado=None):
    """Imagen con recuadros de deteccion (y banner de resultado si se indica)."""
    out = img.copy()
    if m.get("caja_contenedor"):
        x, y, w, h = m["caja_contenedor"]
        cv2.rectangle(out, (x, y), (x + w, y + h), (255, 255, 0), 2)
        cv2.putText(out, "Contenedor", (x + 6, min(y + 20, ALTO - 40)),
                    cv2.FONT_HERSHEY_SIMPLEX, 0.5, (255, 255, 0), 1, cv2.LINE_AA)
    if m.get("caja_tapa"):
        x, y, w, h = m["caja_tapa"]
        cv2.rectangle(out, (x, y), (x + w, y + h), (0, 165, 255), 2)
        cv2.putText(out, "Tapa", (x + 6, max(y - 8, 46)),
                    cv2.FONT_HERSHEY_SIMPLEX, 0.5, (0, 165, 255), 1, cv2.LINE_AA)
    if m.get("linea_tapa"):
        cx, cy, vx, vy = m["linea_tapa"]
        p1 = (int(cx - vx * 220), int(cy - vy * 220))
        p2 = (int(cx + vx * 220), int(cy + vy * 220))
        cv2.line(out, p1, p2, (0, 255, 255), 1, cv2.LINE_AA)
    if m.get("inclinacion") is not None:
        cv2.rectangle(out, (0, ALTO - 30), (ANCHO, ALTO), (0, 0, 0), -1)
        txt = (f"Inclinacion {m['inclinacion']:.1f} deg | Altura {m['altura']:.0f} px | "
               f"Desalineacion {m['desalineacion']:.0f} px")
        cv2.putText(out, _ascii(txt), (10, ALTO - 10), cv2.FONT_HERSHEY_SIMPLEX, 0.55,
                    (255, 255, 255), 1, cv2.LINE_AA)
    if resultado:
        color = (46, 125, 50) if resultado == "APROBADO" else (40, 40, 198)
        cv2.rectangle(out, (0, 0), (ANCHO, 36), color, -1)
        cv2.putText(out, resultado, (12, 26), cv2.FONT_HERSHEY_SIMPLEX, 0.85,
                    (255, 255, 255), 2, cv2.LINE_AA)
    return out


def inspeccionar(nombre, img, lote="", umbrales=None, estacion="", operador=""):
    """Inspecciona una imagen. Devuelve (registro, imagen_con_recuadros_BGR)."""
    t0 = time.perf_counter()
    m = medir(img, umbrales)
    res, causas, just = decidir(m, umbrales)
    ahora = datetime.now()
    registro = {
        "Lote": lote, "Estación": estacion, "Operador": operador, "Imagen": nombre,
        "Fecha": ahora.strftime("%Y-%m-%d"), "Hora": ahora.strftime("%H:%M:%S"),
        "Contenedor": "Presente" if m["contenedor"] else "Ausente",
        "Tapa": "Presente" if m["tapa"] else "Ausente",
        "Inclinación (°)": m["inclinacion"], "Altura (px)": m["altura"],
        "Desalineación (px)": m["desalineacion"],
        "Resultado": res, "Justificación Técnica": just,
        "Causa de rechazo": ", ".join(causas) if causas else "Ninguna",
    }
    vista = dibujar(img, m, res)
    registro["_t_ms"] = round((time.perf_counter() - t0) * 1000, 1)  # solo dashboard
    return registro, vista


# ---------------------------------------------------------------------------
# CALIBRACION ASISTIDA POR EL OPERADOR (aprendizaje supervisado ligero)
# ---------------------------------------------------------------------------
REJILLA_HSV = {"azul_s_min": [100, 120, 150], "azul_v_min": [45, 60, 80],
               "blanco_s_max": [60, 70, 85], "blanco_v_min": [100, 110]}


def leer_etiquetas(ruta):
    """CSV imagen,etiqueta (APROBADO/RECHAZADO) -> dict nombre -> bool."""
    import csv
    et = {}
    with open(ruta, newline="", encoding="utf-8-sig") as f:
        for fila in csv.DictReader(f):
            v = (fila.get("etiqueta") or "").strip().upper()
            if v in ("APROBADO", "RECHAZADO"):
                et[fila["imagen"].strip()] = v == "APROBADO"
    return et


def leer_referencia(ruta):
    """CSV de referencia: imagen, etiqueta, causa_esperada (varias separadas por ';').
    Devuelve dict nombre -> {"apta": bool, "causas": frozenset, "descripcion": str}."""
    import csv
    ref = {}
    with open(ruta, newline="", encoding="utf-8-sig") as f:
        for fila in csv.DictReader(f):
            v = (fila.get("etiqueta") or "").strip().upper()
            if v not in ("APROBADO", "RECHAZADO"):
                continue
            causas = [c.strip() for c in (fila.get("causa_esperada") or "").split(";")
                      if c.strip() and c.strip() != "Ninguna"]
            ref[fila["imagen"].strip()] = {"apta": v == "APROBADO", "causas": frozenset(causas),
                                           "descripcion": (fila.get("descripcion") or "").strip()}
    return ref


def validar(imagenes, referencia, umbrales=None):
    """Compara la decision del sistema con la referencia.
    Devuelve (DataFrame por imagen, dict de metricas)."""
    import pandas as pd
    filas = []
    for nombre, img in imagenes.items():
        if nombre not in referencia:
            continue
        r = referencia[nombre]
        m = medir(img, umbrales)
        res, causas, _ = decidir(m, umbrales)
        filas.append({
            "Imagen": nombre,
            "Esperado": "APROBADO" if r["apta"] else "RECHAZADO",
            "Obtenido": res,
            "Resultado correcto": (res == "APROBADO") == r["apta"],
            "Causas esperadas": "; ".join(sorted(r["causas"])) or "Ninguna",
            "Causas obtenidas": "; ".join(sorted(causas)) or "Ninguna",
            "Causa correcta": frozenset(causas) == r["causas"],
            "Inclinación (°)": m["inclinacion"], "Altura (px)": m["altura"],
            "Desalineación (px)": m["desalineacion"],
            "Descripción de referencia": r["descripcion"]})
    df = pd.DataFrame(filas)
    if df.empty:
        return df, {}
    esp, obt = df["Esperado"] == "APROBADO", df["Obtenido"] == "APROBADO"
    vp, vn = int((esp & obt).sum()), int((~esp & ~obt).sum())
    fp, fn = int((~esp & obt).sum()), int((esp & ~obt).sum())
    div = lambda a, b: round(a / b, 4) if b else None
    met = {"n": len(df), "exactitud": div(vp + vn, len(df)),
           "sensibilidad": div(vp, vp + fn), "especificidad": div(vn, vn + fp),
           "precision": div(vp, vp + fp), "exactitud_causa": round(float(df["Causa correcta"].mean()), 4),
           "VP": vp, "VN": vn, "FP": fp, "FN": fn}
    return df, met


def _caracteristicas(imagenes, u):
    X = np.full((len(imagenes), 3), np.nan)
    presentes = np.zeros(len(imagenes), dtype=bool)
    for i, img in enumerate(imagenes):
        m = medir(img, u)
        if m["contenedor"] and m["tapa"]:
            presentes[i] = True
            X[i] = [m["inclinacion"], m["altura"], m["desalineacion"]]
    return X, presentes


INF = 1e9


def _candidatos(col):
    """Umbrales candidatos = puntos medios entre valores observados consecutivos.
    Devuelve (candidatos, cota_inferior, cota_superior) de cada intervalo."""
    v = np.unique(col[~np.isnan(col)])
    if len(v) == 0:
        return np.array([0.0]), np.array([-INF]), np.array([INF])
    cand = np.concatenate([[v[0] - 0.5], (v[:-1] + v[1:]) / 2, [v[-1] + 0.5]])
    return cand, np.concatenate([[-INF], v]), np.concatenate([v, [INF]])


def _predecir(X, presentes, u):
    Xs = np.nan_to_num(X, nan=1e9)
    return (presentes & (Xs[:, 0] <= u["inclinacion"]) & (Xs[:, 1] <= u["altura"])
            & (Xs[:, 2] <= u["desalineacion"]))


def _exactitud_balanceada(y, pred):
    pos, neg = y.sum(), (~y).sum()
    tpr = (pred & y).sum() / pos if pos else 1.0
    tnr = (~pred & ~y).sum() / neg if neg else 1.0
    return (tpr + tnr) / 2


def buscar_umbrales(X, presentes, y, previos):
    """
    Busqueda exhaustiva (vectorizada) de los 3 umbrales de decision que maximizan
    la exactitud balanceada, con REGULARIZACION contra el sobreajuste:
      - Si el umbral actual es compatible con las respuestas del operador, se CONSERVA.
      - Si no, se coloca a 1/3 del margen factible desde el lado aprobado
        (lado conservador: es peor aceptar una pieza mala que rechazar una buena).
      - Siempre dentro de rangos fisicos plausibles (LIMITES).
    """
    info = [_candidatos(X[presentes, j]) for j in range(3)]
    cands = [i[0] for i in info]
    Xs = np.nan_to_num(X, nan=1e9)
    B = [(Xs[:, j][None, :] <= cands[j][:, None]).astype(float) for j in range(3)]
    pos, neg = max(int(y.sum()), 1), max(int((~y).sum()), 1)
    tp = np.einsum("ai,bi,ci->abc", B[0] * (presentes & y), B[1], B[2])
    fp = np.einsum("ai,bi,ci->abc", B[0] * (presentes & ~y), B[1], B[2])
    score = (tp / pos + (neg - fp) / neg) / 2
    mejor = float(score.max())
    idx = np.argwhere(score >= mejor - 1e-12)
    ties = np.array([[cands[j][i[j]] for j in range(3)] for i in idx])

    prop = []
    for j, k in enumerate(CLAVES_DECISION):
        lo = max(min(info[j][1][i[j]] for i in idx), LIMITES[k][0])   # intervalo factible
        hi = min(max(info[j][2][i[j]] for i in idx), LIMITES[k][1])   # [lo, hi)
        if lo <= previos[k] < hi:
            t = previos[k]
        else:
            t = lo + max(hi - lo, 0.0) / 3
        prop.append(min(max(t, LIMITES[k][0]), LIMITES[k][1]))
    prop = np.array(prop)
    u_prop = dict(zip(CLAVES_DECISION, prop))
    if _exactitud_balanceada(y, _predecir(X, presentes, u_prop)) < mejor - 1e-12:
        rango = np.maximum(ties.max(axis=0) - ties.min(axis=0), 1e-9)   # no era optima
        prop = ties[int(np.argmin((((ties - prop) / rango) ** 2).sum(axis=1)))]
    return {k: round(float(v), 1) for k, v in zip(CLAVES_DECISION, prop)}, mejor


def _tablas_hsv(imgs, combos, base, progreso=None):
    """Caracteristicas de todas las imagenes para todas las combinaciones HSV
    (las mascaras se reutilizan: 15 segmentaciones por imagen en vez de 54)."""
    X = np.full((len(combos), len(imgs), 3), np.nan)
    P = np.zeros((len(combos), len(imgs)), dtype=bool)
    for i, img in enumerate(imgs):
        hsv = _hsv(img)
        azules, blancos = {}, {}
        for ci, c in enumerate(combos):
            u = {**base, **c}
            ka, kb = (u["azul_s_min"], u["azul_v_min"]), (u["blanco_s_max"], u["blanco_v_min"])
            if ka not in azules:
                azules[ka] = _mascara_azul(hsv, u)
            if kb not in blancos:
                blancos[kb] = _mascara_blanco(hsv, u)
            f = _medir_mascaras(azules[ka], blancos[kb])
            if f["contenedor"] and f["tapa"]:
                P[ci, i] = True
                X[ci, i] = [f["inclinacion"], f["altura"], f["desalineacion"]]
        if progreso:
            progreso((i + 1) / len(imgs))
    return X, P


def calibrar(imagenes, etiquetas, calibrar_hsv=True, progreso=None):
    """
    imagenes : dict nombre -> imagen BGR normalizada
    etiquetas: dict nombre -> True (APTA) / False (RECHAZADA)  [respuestas del operador]
    """
    nombres = [n for n in imagenes if n in etiquetas]
    y = np.array([etiquetas[n] for n in nombres], dtype=bool)
    if len(nombres) < 6 or y.all() or not y.any():
        raise ValueError("Se necesitan al menos 6 respuestas, con piezas APTAS y RECHAZADAS.")
    imgs = [imagenes[n] for n in nombres]
    base = dict(UMBRALES_DEFECTO)

    combos = [{}]
    if calibrar_hsv:
        claves = list(REJILLA_HSV)
        combos = [dict(zip(claves, v)) for v in itertools.product(*REJILLA_HSV.values())]
    Xc, Pc = _tablas_hsv(imgs, combos, base, progreso)
    mejor = None
    for ci, c in enumerate(combos):
        t, sc = buscar_umbrales(Xc[ci], Pc[ci], y, base)
        dist = sum(abs(c[k] - base[k]) / base[k] for k in c)      # preferir HSV por defecto
        clave = (sc, -dist)
        if mejor is None or clave > mejor[0]:
            mejor = (clave, {**base, **c, **t}, Xc[ci], Pc[ci])
    _, u_fin, X, pres = mejor

    pred = _predecir(X, pres, u_fin)
    tp, tn = int((pred & y).sum()), int((~pred & ~y).sum())
    fp, fn = int((pred & ~y).sum()), int((~pred & y).sum())

    # exactitud con los umbrales por defecto (referencia de comparacion)
    X0, p0 = _caracteristicas(imgs, base)
    exact_defecto = float((_predecir(X0, p0, base) == y).mean())

    # validacion cruzada leave-one-out (HSV fijo): cada pieza se predice sin usarla
    aciertos = 0
    for k in range(len(y)):
        m = np.ones(len(y), dtype=bool)
        m[k] = False
        if y[m].all() or not y[m].any():
            t = {kk: base[kk] for kk in CLAVES_DECISION}
        else:
            t, _ = buscar_umbrales(X[m], pres[m], y[m], base)
        aciertos += int(_predecir(X[k:k + 1], pres[k:k + 1], {**u_fin, **t})[0] == y[k])

    return {
        "fecha": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        "n_respuestas": len(y), "n_aptas": int(y.sum()), "n_rechazadas": int((~y).sum()),
        "umbrales": {k: (round(float(v), 1) if isinstance(v, (float, np.floating)) else int(v))
                     for k, v in u_fin.items()},
        "exactitud": round((tp + tn) / len(y), 4),
        "exactitud_defecto": round(exact_defecto, 4),
        "exactitud_loocv": round(aciertos / len(y), 4),
        "matriz_confusion": {"APTAS bien clasificadas": tp, "RECHAZADAS bien clasificadas": tn,
                             "RECHAZADAS aceptadas por error": fp,
                             "APTAS rechazadas por error": fn},
        "errores": [nombres[i] for i in range(len(y)) if pred[i] != y[i]],
    }


def guardar_calibracion(resultado, ruta=RUTA_CALIBRACION):
    Path(ruta).write_text(json.dumps(resultado, indent=2, ensure_ascii=False), encoding="utf-8")
    UMBRALES.update(resultado["umbrales"])


def restablecer_calibracion(ruta=RUTA_CALIBRACION):
    p = Path(ruta)
    if p.exists():
        p.unlink()
    UMBRALES.clear()
    UMBRALES.update(UMBRALES_DEFECTO)


# ---------------------------------------------------------------------------
# EXCEL: base de datos historica (acumula, nunca sobrescribe)
# ---------------------------------------------------------------------------
def _nuevo_libro():
    wb = Workbook()
    ws = wb.active
    ws.title = "Historico"
    ws.append(COLUMNAS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF", name="Arial")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    anchos = [18, 16, 16, 20, 12, 10, 12, 10, 14, 12, 16, 13, 95, 30]
    for i, a in enumerate(anchos, 1):
        ws.column_dimensions[get_column_letter(i)].width = a
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(COLUMNAS))}1"
    return wb


def guardar_excel(ruta, registros):
    """Agrega los registros al historico (crea el archivo si no existe)."""
    ruta = Path(ruta)
    if ruta.exists():
        wb = load_workbook(ruta)
        ws = wb["Historico"] if "Historico" in wb.sheetnames else None
        if ws is None or [c.value for c in ws[1]] != COLUMNAS:
            # archivo de una version anterior con otro formato: se conserva aparte
            wb.close()
            ruta.rename(ruta.with_name(f"{ruta.stem}_anterior_{datetime.now():%Y%m%d_%H%M%S}"
                                       f"{ruta.suffix}"))
            wb = _nuevo_libro()
            ws = wb["Historico"]
    else:
        wb = _nuevo_libro()
        ws = wb["Historico"]

    col_res = COLUMNAS.index("Resultado") + 1
    for reg in registros:
        ws.append([reg[c] for c in COLUMNAS])
        ok = reg["Resultado"] == "APROBADO"
        celda = ws.cell(ws.max_row, col_res)
        celda.font = Font(bold=True, color="006100" if ok else "9C0006")
        celda.fill = PatternFill("solid", fgColor="C6EFCE" if ok else "FFC7CE")

    if "Resumen" in wb.sheetnames:
        del wb["Resumen"]
    rs = wb.create_sheet("Resumen")
    L = get_column_letter(col_res)
    rs.append(["Indicador", "Valor"])
    rs.append(["Total inspecciones", "=COUNTA(Historico!A2:A100000)"])
    rs.append(["Aprobados", f'=COUNTIF(Historico!{L}2:{L}100000,"APROBADO")'])
    rs.append(["Rechazados", f'=COUNTIF(Historico!{L}2:{L}100000,"RECHAZADO")'])
    rs.append(["% Aprobados", "=IF(B2=0,0,B3/B2)"])
    rs.append(["% Rechazados", "=IF(B2=0,0,B4/B2)"])
    rs["B5"].number_format = rs["B6"].number_format = "0.0%"
    for c in rs[1]:
        c.font = Font(bold=True, color="FFFFFF", name="Arial")
        c.fill = PatternFill("solid", fgColor="1F4E78")
    rs.column_dimensions["A"].width = 24
    rs.column_dimensions["B"].width = 14
    wb.save(ruta)
